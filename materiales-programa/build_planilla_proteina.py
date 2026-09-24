"""Genera la planilla de proteína y fibra del programa Diseña tu Jubilación.

Salida: materiales-programa/planilla-proteina-fibra.xlsx

Tres hojas para el participante (Mi objetivo, Mi día, Alimentos) y una de
fuentes. Todo se calcula con fórmulas, así que funciona en Excel y en Google
Sheets sin macros. Los valores de referencia están etiquetados por nivel de
evidencia, igual que la tabla de Medicina preventiva: G = guía clínica,
C = consenso de sociedad, E = opinión de experto.

Borrador de Claude (24/09/2026). Antes de entregarla a participantes la
revisa el médico aliado o una nutricionista.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).with_name("planilla-proteina-fibra.xlsx")

VERDE = "0E3A2F"
CARBON = "1A2421"
AMBAR = "E65F2B"
CREMA = "F6E7B0"
CLARO = "F8F9FA"

F_TITULO = Font(name="Lora", size=18, bold=True, color=VERDE)
F_SUB = Font(name="Poppins", size=10, color=CARBON, italic=True)
F_HEAD = Font(name="Poppins", size=10, bold=True, color=CLARO)
F_LABEL = Font(name="Poppins", size=10, bold=True, color=CARBON)
F_TXT = Font(name="Poppins", size=10, color=CARBON)
F_BIG = Font(name="Lora", size=14, bold=True, color=VERDE)
FILL_HEAD = PatternFill("solid", fgColor=VERDE)
FILL_INPUT = PatternFill("solid", fgColor="FFF6D6")
FILL_OUT = PatternFill("solid", fgColor="EAF1EE")
LINEA = Side(style="thin", color="C9D3CF")
BORDE = Border(left=LINEA, right=LINEA, top=LINEA, bottom=LINEA)
WRAP = Alignment(wrap_text=True, vertical="top")

# Por 100 g de alimento listo para comer: proteína, fibra, grasa,
# carbohidratos (g) y kcal, más una porción habitual en gramos.
# Valores aproximados de USDA FoodData Central. Los marcados con * son
# alimentos chilenos estimados por analogía: verificar con la tabla chilena.
ALIMENTOS = [
    # nombre, porción g, descripción porción, prot, fibra, grasa, carb, kcal
    ("Huevo entero", 50, "1 unidad", 12.6, 0, 9.5, 0.7, 143),
    ("Clara de huevo", 33, "1 clara", 10.9, 0, 0.2, 0.7, 52),
    ("Pechuga de pollo, cocida", 120, "1 filete", 31, 0, 3.6, 0, 165),
    ("Trutro de pollo sin piel, cocido", 120, "1 trutro", 26, 0, 8, 0, 180),
    ("Vacuno magro (posta), cocido", 120, "1 bistec", 29, 0, 7, 0, 185),
    ("Carne molida 10% grasa, cocida", 120, "1 porción", 26, 0, 11, 0, 217),
    ("Lomo de cerdo, cocido", 120, "1 chuleta", 27, 0, 6, 0, 170),
    ("Salmón, cocido", 120, "1 filete", 25, 0, 12, 0, 206),
    ("Merluza, cocida", 120, "1 filete", 20, 0, 1.5, 0, 95),
    ("Atún en agua, escurrido", 80, "1 lata chica", 25, 0, 1, 0, 116),
    ("Jurel en lata, escurrido *", 80, "1/2 lata", 21, 0, 8, 0, 160),
    ("Jamón de pavo", 30, "2 láminas", 17, 0, 2, 2, 100),
    ("Queso fresco *", 30, "1 rebanada", 17, 0, 13, 3, 200),
    ("Quesillo *", 50, "1 rebanada gruesa", 11, 0, 8, 3, 130),
    ("Queso mantecoso o gauda", 20, "1 lámina", 25, 0, 27, 2, 356),
    ("Queso cottage", 100, "1/2 taza", 11, 0, 4.3, 3.4, 98),
    ("Yogur natural", 150, "1 pote", 3.5, 0, 3.3, 4.7, 61),
    ("Yogur griego natural descremado", 150, "1 pote", 10, 0, 0.4, 3.6, 59),
    ("Leche semidescremada", 200, "1 taza", 3.3, 0, 1.5, 4.8, 46),
    ("Bebida de soya sin azúcar", 200, "1 taza", 3.3, 0.6, 1.8, 3, 43),
    ("Proteína de suero en polvo", 30, "1 medida", 78, 0, 5, 8, 380),
    ("Lentejas, cocidas", 150, "3/4 taza", 9, 7.9, 0.4, 20, 116),
    ("Porotos negros, cocidos", 150, "3/4 taza", 8.9, 8.7, 0.5, 23.7, 132),
    ("Porotos bayos o coscorrón, cocidos *", 150, "3/4 taza", 9, 9, 0.7, 26, 143),
    ("Garbanzos, cocidos", 150, "3/4 taza", 8.9, 7.6, 2.6, 27, 164),
    ("Arvejas, cocidas", 100, "1/2 taza", 5.4, 5.5, 0.2, 15.6, 84),
    ("Tofu firme", 100, "1 trozo", 17, 2, 9, 3, 144),
    ("Edamame", 100, "2/3 taza", 11.9, 5.2, 5.2, 8.9, 121),
    ("Hummus", 60, "4 cucharadas", 7.9, 6, 9.6, 14, 166),
    ("Quinoa, cocida", 150, "3/4 taza", 4.4, 2.8, 1.9, 21.3, 120),
    ("Arroz blanco, cocido", 150, "3/4 taza", 2.7, 0.4, 0.3, 28, 130),
    ("Arroz integral, cocido", 150, "3/4 taza", 2.6, 1.8, 0.9, 23, 112),
    ("Fideos, cocidos", 150, "1 taza", 5.8, 1.8, 0.9, 31, 158),
    ("Fideos integrales, cocidos", 150, "1 taza", 5.8, 4.5, 0.9, 30, 149),
    ("Pan blanco (marraqueta o hallulla) *", 100, "1 unidad", 9, 2.7, 3.2, 49, 265),
    ("Pan integral", 60, "2 rebanadas", 13, 7, 3.4, 41, 247),
    ("Avena", 40, "4 cucharadas", 13, 10, 7, 66, 389),
    ("Papa, cocida", 150, "1 mediana", 1.9, 1.8, 0.1, 20, 87),
    ("Choclo, cocido", 100, "1/2 taza", 3.4, 2.4, 1.5, 21, 96),
    ("Brócoli, cocido", 100, "1 taza", 2.4, 3.3, 0.4, 7.2, 35),
    ("Espinaca, cocida", 100, "1/2 taza", 3, 2.4, 0.3, 3.8, 23),
    ("Lechuga", 50, "1 plato", 1.4, 1.3, 0.2, 2.9, 15),
    ("Tomate", 120, "1 mediano", 0.9, 1.2, 0.2, 3.9, 18),
    ("Zanahoria, cruda", 80, "1 mediana", 0.9, 2.8, 0.2, 9.6, 41),
    ("Palta", 50, "1/4 unidad", 2, 6.7, 14.7, 8.5, 160),
    ("Manzana", 150, "1 mediana", 0.3, 2.4, 0.2, 13.8, 52),
    ("Plátano", 120, "1 mediano", 1.1, 2.6, 0.3, 22.8, 89),
    ("Naranja", 150, "1 mediana", 0.9, 2.4, 0.1, 11.8, 47),
    ("Pera", 150, "1 mediana", 0.4, 3.1, 0.1, 15.2, 57),
    ("Kiwi", 75, "1 unidad", 1.1, 3, 0.5, 14.7, 61),
    ("Frutillas", 150, "1 taza", 0.7, 2, 0.3, 7.7, 32),
    ("Arándanos", 100, "2/3 taza", 0.7, 2.4, 0.3, 14.5, 57),
    ("Nueces", 30, "1 puñado", 15, 6.7, 65, 14, 654),
    ("Almendras", 30, "1 puñado", 21, 12.5, 50, 22, 579),
    ("Maní", 30, "1 puñado", 26, 8.5, 49, 16, 567),
    ("Mantequilla de maní", 16, "1 cucharada", 25, 6, 50, 20, 588),
    ("Semillas de chía", 15, "1 cucharada", 17, 34, 31, 42, 486),
    ("Linaza molida", 10, "1 cucharada", 18, 27, 42, 29, 534),
    ("Aceite de oliva", 10, "1 cucharada", 0, 0, 100, 0, 884),
]

COMIDAS = ["Desayuno", "Colación", "Almuerzo", "Once", "Cena"]
FILAS_POR_COMIDA = 5


def encabezado(ws, fila, valores, anchos=None):
    for i, v in enumerate(valores, start=1):
        c = ws.cell(row=fila, column=i, value=v)
        c.font = F_HEAD
        c.fill = FILL_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDE
    if anchos:
        for i, w in enumerate(anchos, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w


def hoja_alimentos(wb):
    ws = wb.create_sheet("Alimentos")
    ws["A1"] = "Alimentos"
    ws["A1"].font = F_TITULO
    ws["A2"] = ("Valores por 100 g de alimento listo para comer. Los marcados con * son "
                "estimaciones para alimentos chilenos. Puedes agregar los tuyos en las filas vacías.")
    ws["A2"].font = F_SUB
    encabezado(ws, 4, ["Alimento", "Porción (g)", "Porción habitual", "Proteína", "Fibra",
                       "Grasa", "Carbohidratos", "kcal"],
               [38, 11, 18, 10, 9, 9, 13, 9])
    for r, fila in enumerate(ALIMENTOS, start=5):
        for c, v in enumerate(fila, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = F_TXT
            cell.border = BORDE
    ultima = 4 + len(ALIMENTOS) + 20  # 20 filas libres para alimentos propios
    for r in range(5 + len(ALIMENTOS), ultima + 1):
        for c in range(1, 9):
            cell = ws.cell(row=r, column=c)
            cell.fill = FILL_INPUT
            cell.border = BORDE
    ws.freeze_panes = "A5"
    return ws, ultima


def hoja_objetivo(wb):
    ws = wb.active
    ws.title = "Mi objetivo"
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 60

    ws["A1"] = "Cuánta proteína y fibra necesitas"
    ws["A1"].font = F_TITULO
    ws["A2"] = ("Completa las celdas amarillas. La planilla calcula tu objetivo diario; "
                "en la hoja Mi día puedes armar un día de comidas y ver si llegas.")
    ws["A2"].font = F_SUB
    ws.merge_cells("A2:D2")

    datos = [
        ("Sexo", "Mujer"),
        ("Edad (años)", 60),
        ("Peso (kg)", 70),
        ("Estatura (cm)", 165),
        ("Actividad física", "Moderada (3 a 5 días a la semana)"),
        ("Perfil de proteína", "Mayor de 60, activo o entrenando fuerza (1,2 g/kg)"),
    ]
    encabezado(ws, 4, ["Tus datos", "", "", ""])
    for i, (label, val) in enumerate(datos, start=5):
        ws.cell(row=i, column=1, value=label).font = F_LABEL
        c = ws.cell(row=i, column=2, value=val)
        c.fill = FILL_INPUT
        c.font = F_TXT
        c.border = BORDE
    ws.merge_cells("B9:D9")
    ws.merge_cells("B10:D10")

    dv_sexo = DataValidation(type="list", formula1='"Mujer,Hombre"', allow_blank=False)
    dv_act = DataValidation(type="list", formula1="=Listas!$A$2:$A$5", allow_blank=False)
    dv_perfil = DataValidation(type="list", formula1="=Listas!$C$2:$C$6", allow_blank=False)
    for dv, ref in ((dv_sexo, "B5"), (dv_act, "B9"), (dv_perfil, "B10")):
        ws.add_data_validation(dv)
        dv.add(ref)

    encabezado(ws, 12, ["Tu objetivo diario", "Valor", "Nivel", "De dónde sale"])
    filas = [
        ("Índice de masa corporal",
         "=ROUND(B7/(B8/100)^2,1)", "",
         "Si es 30 o más, el cálculo por kilo sobreestima: conversa con tu médico o nutricionista qué peso usar."),
        ("Energía estimada (kcal al día)",
         '=ROUND((10*B7+6.25*B8-5*B6+IF(B5="Hombre",5,-161))*VLOOKUP(B9,Listas!$A$2:$B$5,2,FALSE),0)',
         "C", "Ecuación de Mifflin-St Jeor por el factor de actividad. Es una estimación, no una medición."),
        ("Proteína: objetivo del día (g)",
         "=ROUND(B7*VLOOKUP(B10,Listas!$C$2:$E$6,2,FALSE),0)",
         "=VLOOKUP(B10,Listas!$C$2:$E$6,3,FALSE)",
         "Según el perfil que elegiste. Ver la tabla de perfiles más abajo."),
        ("Proteína: mínimo oficial (g)", "=ROUND(B7*0.8,0)", "G",
         "0,8 g por kilo: ingesta recomendada para adultos (IOM). Evita la carencia, no es un óptimo para mayores."),
        ("Proteína por comida principal (g)", "=ROUND(MAX(25,B15/3),0)", "C",
         "Repartida en tres comidas de al menos 25 a 30 g, en vez de concentrarla en una (PROT-AGE, ESPEN)."),
        ("Fibra: objetivo del día (g)", "=ROUND(B14/1000*14,0)", "G",
         "14 g por cada 1.000 kcal (IOM, base de las guías dietarias de EE. UU.)."),
        ("Fibra: referencia por edad y sexo (g)",
         '=IF(B5="Hombre",IF(B6>50,30,38),IF(B6>50,21,25))', "G",
         "Ingesta adecuada del IOM para tu edad y sexo. La OMS pide al menos 25 g al día a todos los adultos."),
    ]
    for i, (label, formula, nivel, fuente) in enumerate(filas, start=13):
        ws.cell(row=i, column=1, value=label).font = F_LABEL
        v = ws.cell(row=i, column=2, value=formula)
        v.fill = FILL_OUT
        v.font = F_BIG if "objetivo" in label else F_TXT
        v.border = BORDE
        ws.cell(row=i, column=3, value=nivel).font = F_TXT
        f = ws.cell(row=i, column=4, value=fuente)
        f.font = F_TXT
        f.alignment = WRAP
        ws.row_dimensions[i].height = 32

    encabezado(ws, 21, ["Perfiles de proteína", "g por kilo", "Nivel", "Fuente"])
    for i, r in enumerate(range(2, 7), start=22):
        ws.cell(row=i, column=1, value=f"=Listas!C{r}").font = F_TXT
        ws.cell(row=i, column=2, value=f"=Listas!D{r}").font = F_TXT
        ws.cell(row=i, column=3, value=f"=Listas!E{r}").font = F_TXT
        c = ws.cell(row=i, column=4, value=f"=Listas!F{r}")
        c.font = F_TXT
        c.alignment = WRAP
        ws.row_dimensions[i].height = 30

    ws["A28"] = "Antes de subir tu proteína"
    ws["A28"].font = Font(name="Lora", size=12, bold=True, color=AMBAR)
    avisos = [
        "Si tienes enfermedad renal crónica o te han dicho que tu función renal está baja, no subas la proteína sin hablarlo con tu médico.",
        "Esta planilla es una herramienta para ordenarte, no una indicación médica ni nutricional. Llévala a tu médico o nutricionista si tienes dudas.",
        "Nivel: G = guía clínica, C = consenso de sociedad, E = opinión de experto. Donde dice E, no hay consenso de guía.",
    ]
    for i, t in enumerate(avisos, start=29):
        c = ws.cell(row=i, column=1, value=t)
        c.font = F_TXT
        c.alignment = WRAP
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=4)
        ws.row_dimensions[i].height = 30
    return ws


def hoja_listas(wb):
    ws = wb.create_sheet("Listas")
    ws["A1"], ws["B1"] = "Actividad", "Factor"
    actividades = [
        ("Sedentaria (casi nada de ejercicio)", 1.2),
        ("Liviana (1 a 3 días a la semana)", 1.375),
        ("Moderada (3 a 5 días a la semana)", 1.55),
        ("Alta (6 a 7 días a la semana)", 1.725),
    ]
    for i, (a, f) in enumerate(actividades, start=2):
        ws.cell(row=i, column=1, value=a)
        ws.cell(row=i, column=2, value=f)
    ws["C1"], ws["D1"], ws["E1"], ws["F1"] = "Perfil", "g/kg", "Nivel", "Fuente"
    perfiles = [
        ("Mínimo oficial para adultos (0,8 g/kg)", 0.8, "G", "IOM: ingesta recomendada (RDA) para adultos sanos."),
        ("Mayor de 60, sano (1,0 g/kg)", 1.0, "C", "PROT-AGE (2013) y ESPEN (2014): 1,0 a 1,2 g/kg en mayores de 65."),
        ("Mayor de 60, activo o entrenando fuerza (1,2 g/kg)", 1.2, "C", "PROT-AGE: 1,2 g/kg o más si hace ejercicio."),
        ("Enfermedad aguda o crónica, con indicación médica (1,5 g/kg)", 1.5, "C", "ESPEN: 1,2 a 1,5 g/kg. Solo con indicación médica."),
        ("Enfoque longevidad, entrenando fuerza (1,6 g/kg)", 1.6, "E", "Opinión de expertos en longevidad; sin consenso de guía."),
    ]
    for i, fila in enumerate(perfiles, start=2):
        for c, v in enumerate(fila, start=3):
            ws.cell(row=i, column=c, value=v)
    ws.sheet_state = "hidden"
    return ws


def hoja_dia(wb, ultima_alimento):
    ws = wb.create_sheet("Mi día", 1)
    anchos = [12, 38, 11, 11, 10, 10, 10, 13, 9]
    ws["A1"] = "Mi día de comidas"
    ws["A1"].font = F_TITULO
    ws["A2"] = ("Elige un alimento en la columna B y escribe los gramos. Si dejas los gramos vacíos, "
                "se usa la porción habitual.")
    ws["A2"].font = F_SUB
    encabezado(ws, 4, ["Comida", "Alimento", "Gramos", "Proteína", "Fibra", "Grasa",
                       "Carbohidratos", "kcal", "Porción"], anchos)
    rango = f"Alimentos!$A$5:$H${ultima_alimento}"
    dv = DataValidation(type="list", formula1=f"=Alimentos!$A$5:$A${ultima_alimento}", allow_blank=True)
    ws.add_data_validation(dv)

    fila = 5
    totales_comida = []
    for comida in COMIDAS:
        inicio = fila
        for k in range(FILAS_POR_COMIDA):
            ws.cell(row=fila, column=1, value=comida if k == 0 else None).font = F_LABEL
            b = ws.cell(row=fila, column=2)
            b.fill = FILL_INPUT
            b.border = BORDE
            dv.add(b.coordinate)
            g = ws.cell(row=fila, column=3)
            g.fill = FILL_INPUT
            g.border = BORDE
            gramos = f'IF(C{fila}="",VLOOKUP(B{fila},{rango},2,FALSE),C{fila})'
            for col, idx in zip("DEFGH", (4, 5, 6, 7, 8)):
                ws[f"{col}{fila}"] = (f'=IF(B{fila}="","",ROUND({gramos}/100*'
                                      f'VLOOKUP(B{fila},{rango},{idx},FALSE),1))')
                ws[f"{col}{fila}"].font = F_TXT
                ws[f"{col}{fila}"].border = BORDE
            ws[f"I{fila}"] = f'=IF(B{fila}="","",VLOOKUP(B{fila},{rango},3,FALSE))'
            ws[f"I{fila}"].font = F_SUB
            fila += 1
        ws.cell(row=fila, column=2, value=f"Total {comida.lower()}").font = F_LABEL
        for col in "DEFGH":
            c = ws[f"{col}{fila}"]
            c.value = f"=SUM({col}{inicio}:{col}{fila - 1})"
            c.font = F_LABEL
            c.fill = FILL_OUT
        totales_comida.append(fila)
        fila += 2

    # Resumen del día
    r = fila
    encabezado_cells = ["", "Resumen del día", "", "Proteína", "Fibra", "Grasa", "Carbohidratos", "kcal"]
    for i, v in enumerate(encabezado_cells, start=1):
        c = ws.cell(row=r, column=i, value=v or None)
        if v:
            c.font = F_HEAD
            c.fill = FILL_HEAD
    ws.cell(row=r + 1, column=2, value="Total del día").font = F_LABEL
    for col in "DEFGH":
        ws[f"{col}{r + 1}"] = "=" + "+".join(f"{col}{t}" for t in totales_comida)
        ws[f"{col}{r + 1}"].font = F_BIG
    ws.cell(row=r + 2, column=2, value="Tu objetivo").font = F_LABEL
    ws[f"D{r + 2}"] = "='Mi objetivo'!B15"
    ws[f"E{r + 2}"] = "='Mi objetivo'!B18"
    ws[f"H{r + 2}"] = "='Mi objetivo'!B14"
    ws.cell(row=r + 3, column=2, value="Cuánto llevas del objetivo").font = F_LABEL
    for col in "DEH":
        c = ws[f"{col}{r + 3}"]
        c.value = f'=IF({col}{r + 2}=0,"",{col}{r + 1}/{col}{r + 2})'
        c.number_format = "0%"
        c.font = F_BIG
    verde = PatternFill("solid", fgColor="CFE6DA")
    ambar = PatternFill("solid", fgColor="FBD9C9")
    for col in "DE":
        ref = f"{col}{r + 3}"
        ws.conditional_formatting.add(ref, CellIsRule(operator="greaterThanOrEqual", formula=["0.9"], fill=verde))
        ws.conditional_formatting.add(ref, CellIsRule(operator="lessThan", formula=["0.9"], fill=ambar))

    r2 = r + 5
    ws.cell(row=r2, column=2, value="Comidas principales con al menos 25 g de proteína").font = F_LABEL
    principales = [totales_comida[0], totales_comida[2], totales_comida[4]]
    ws[f"D{r2}"] = "=" + "+".join(f'IF(D{t}>=25,1,0)' for t in principales)
    ws[f"E{r2}"] = "de 3"
    ws[f"D{r2}"].font = F_BIG
    ws.freeze_panes = "A5"
    return ws


def hoja_fuentes(wb):
    ws = wb.create_sheet("Fuentes")
    ws.column_dimensions["A"].width = 110
    lineas = [
        ("Fuentes", F_TITULO),
        ("Borrador de Claude para Diseña tu Jubilación, 24/09/2026. Lo revisa el médico aliado o una nutricionista antes de entregarlo.", F_SUB),
        ("", F_TXT),
        ("Proteína", F_LABEL),
        ("Institute of Medicine (IOM). Dietary Reference Intakes for Energy, Carbohydrate, Fiber, Fat, Fatty Acids, Cholesterol, Protein, and Amino Acids. 2005. RDA 0,8 g/kg.", F_TXT),
        ("Bauer J. y col. Evidence-based recommendations for optimal dietary protein intake in older people: a position paper from the PROT-AGE Study Group. J Am Med Dir Assoc. 2013.", F_TXT),
        ("Deutz NE y col. Protein intake and exercise for optimal muscle function with aging: recommendations from the ESPEN Expert Group. Clin Nutr. 2014.", F_TXT),
        ("Reparto por comida (25 a 30 g): PROT-AGE y ESPEN, a partir de estudios de síntesis de proteína muscular. Nivel C.", F_TXT),
        ("1,6 g/kg: opinión de expertos en longevidad y metaanálisis de entrenamiento de fuerza (Morton y col., Br J Sports Med 2018) en adultos más jóvenes. Nivel E para este público.", F_TXT),
        ("", F_TXT),
        ("Fibra", F_LABEL),
        ("IOM 2005: ingesta adecuada de 14 g por 1.000 kcal; 30 g (hombres) y 21 g (mujeres) desde los 51 años.", F_TXT),
        ("OMS: al menos 25 g de fibra al día en adultos (Healthy diet, nota descriptiva).", F_TXT),
        ("", F_TXT),
        ("Energía", F_LABEL),
        ("Mifflin MD, St Jeor ST y col. A new predictive equation for resting energy expenditure in healthy individuals. Am J Clin Nutr. 1990.", F_TXT),
        ("", F_TXT),
        ("Alimentos", F_LABEL),
        ("USDA FoodData Central (valores aproximados, redondeados). Los marcados con * son estimaciones para alimentos chilenos: verificar con la tabla de composición de alimentos chilena antes de entregar.", F_TXT),
    ]
    for i, (t, f) in enumerate(lineas, start=1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = f
        c.alignment = WRAP
    return ws


def main():
    wb = Workbook()
    hoja_objetivo(wb)
    _, ultima = hoja_alimentos(wb)
    hoja_dia(wb, ultima)
    hoja_fuentes(wb)
    hoja_listas(wb)
    wb.active = 0
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
